"""SQLite tracker with Excel export."""
import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from .models import Job

SCHEMA = """
CREATE TABLE IF NOT EXISTS jobs (
    id TEXT PRIMARY KEY,
    title TEXT, company TEXT, location TEXT, url TEXT,
    description TEXT, source TEXT, posted_at TEXT, salary TEXT,
    score INTEGER DEFAULT 0, score_reason TEXT DEFAULT '',
    status TEXT DEFAULT 'new', fetched_at TEXT,
    tailored_path TEXT DEFAULT '', applied_at TEXT DEFAULT '', notes TEXT DEFAULT '',
    fuzzy_key TEXT DEFAULT '',
    experience_level TEXT DEFAULT '', company_careers_url TEXT DEFAULT '',
    scan_id INTEGER DEFAULT 0,
    apply_type TEXT DEFAULT 'apply_now'
);

CREATE TABLE IF NOT EXISTS scan_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    mode TEXT DEFAULT 'full',
    sources TEXT DEFAULT '',
    started_at TEXT DEFAULT '',
    finished_at TEXT DEFAULT '',
    status TEXT DEFAULT 'running',
    fetched_count INTEGER DEFAULT 0,
    unique_count INTEGER DEFAULT 0,
    new_count INTEGER DEFAULT 0,
    shortlisted_count INTEGER DEFAULT 0,
    error TEXT DEFAULT '',
    new_job_ids TEXT DEFAULT '[]'
);
"""


class Tracker:
    def __init__(self, db_path: Path):
        db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(db_path)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(SCHEMA)
        self._migrate()

    def _migrate(self):
        """Add columns introduced after the initial schema to older DB files."""
        cols = {row["name"] for row in self.conn.execute("PRAGMA table_info(jobs)")}
        if "fuzzy_key" not in cols:
            self.conn.execute("ALTER TABLE jobs ADD COLUMN fuzzy_key TEXT DEFAULT ''")
            self.conn.commit()
        if "experience_level" not in cols:
            self.conn.execute("ALTER TABLE jobs ADD COLUMN experience_level TEXT DEFAULT ''")
            self.conn.commit()
        if "company_careers_url" not in cols:
            self.conn.execute("ALTER TABLE jobs ADD COLUMN company_careers_url TEXT DEFAULT ''")
            self.conn.commit()
        if "scan_id" not in cols:
            self.conn.execute("ALTER TABLE jobs ADD COLUMN scan_id INTEGER DEFAULT 0")
            self.conn.commit()
        if "apply_type" not in cols:
            self.conn.execute("ALTER TABLE jobs ADD COLUMN apply_type TEXT DEFAULT ''")
            self.conn.commit()

        # Migrate existing rows to accurate apply_type
        self.conn.execute("""
            UPDATE jobs
            SET apply_type = 'quick_apply'
            WHERE LOWER(description) LIKE '%quick apply%' OR LOWER(description) LIKE '%quickapply%' OR LOWER(title) LIKE '%quick apply%'
        """)
        self.conn.execute("""
            UPDATE jobs
            SET apply_type = 'apply_now'
            WHERE apply_type IS NULL OR apply_type = '' OR (LOWER(description) NOT LIKE '%quick apply%' AND LOWER(description) NOT LIKE '%quickapply%' AND LOWER(title) NOT LIKE '%quick apply%')
        """)

        self.conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_fuzzy_key ON jobs(fuzzy_key)")
        self.conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_scan_id ON jobs(scan_id)")
        self.conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_apply_type ON jobs(apply_type)")
        self.conn.commit()

    def is_archived_or_applied(self, job: Job) -> bool:
        """Check if this job (by id, url, or fuzzy_key) was ever marked as applied, archived, or not interested."""
        if self.conn.execute("SELECT 1 FROM jobs WHERE id=? AND status IN ('applied', 'archived', 'not_interested', 'skipped')", (job.id,)).fetchone():
            return True
        if job.fuzzy_key and self.conn.execute("SELECT 1 FROM jobs WHERE fuzzy_key=? AND status IN ('applied', 'archived', 'not_interested', 'skipped')", (job.fuzzy_key,)).fetchone():
            return True
        if job.url and self.conn.execute("SELECT 1 FROM jobs WHERE url=? AND url!='' AND status IN ('applied', 'archived', 'not_interested', 'skipped')", (job.url,)).fetchone():
            return True
        return False

    def is_tracked(self, job: Job) -> bool:
        """Check if this job (by id, url, or fuzzy_key) is already tracked under any status in the database."""
        if self.conn.execute("SELECT 1 FROM jobs WHERE id=?", (job.id,)).fetchone():
            return True
        if job.fuzzy_key and self.conn.execute("SELECT 1 FROM jobs WHERE fuzzy_key=? AND fuzzy_key!=''", (job.fuzzy_key,)).fetchone():
            return True
        if job.url and self.conn.execute("SELECT 1 FROM jobs WHERE url=? AND url!=''", (job.url,)).fetchone():
            return True
        return False

    def upsert(self, job: Job, scan_id: int = 0) -> bool:
        """Insert job; returns True if new, False if already tracked
        (exact URL match or a fuzzy title+company match from another source)."""
        if self.is_tracked(job):
            return False
        d = job.to_dict()
        d["scan_id"] = scan_id
        if not d.get("apply_type"):
            low_desc = (d.get("description") or "").lower()
            d["apply_type"] = "quick_apply" if ("quick apply" in low_desc or "quickapply" in low_desc) else "apply_now"

        self.conn.execute(
            """INSERT INTO jobs (id,title,company,location,url,description,source,
               posted_at,salary,score,score_reason,status,fetched_at,fuzzy_key,
               experience_level,company_careers_url,scan_id,apply_type)
               VALUES (:id,:title,:company,:location,:url,:description,:source,
               :posted_at,:salary,:score,:score_reason,:status,:fetched_at,:fuzzy_key,
               :experience_level,:company_careers_url,:scan_id,:apply_type)""", d)
        self.conn.commit()
        return True

    def create_scan(self, mode: str = "full", sources: str = "") -> int:
        now = datetime.now(timezone.utc).isoformat()
        cur = self.conn.execute(
            """INSERT INTO scan_history (mode, sources, started_at, status, new_job_ids)
               VALUES (?, ?, ?, 'running', '[]')""",
            (mode, sources, now)
        )
        self.conn.commit()
        return cur.lastrowid

    def update_scan(self, scan_id: int, status: str = "completed", fetched_count: int = 0,
                    unique_count: int = 0, new_count: int = 0, shortlisted_count: int = 0,
                    new_job_ids: list[str] | None = None, error: str = ""):
        now = datetime.now(timezone.utc).isoformat()
        job_ids_json = json.dumps(new_job_ids or [])
        self.conn.execute(
            """UPDATE scan_history
               SET finished_at=?, status=?, fetched_count=?, unique_count=?,
                   new_count=?, shortlisted_count=?, new_job_ids=?, error=?
               WHERE id=?""",
            (now, status, fetched_count, unique_count, new_count, shortlisted_count, job_ids_json, error, scan_id)
        )
        self.conn.commit()

    def get_scan_history(self, limit: int = 100, mode: str = "") -> list[sqlite3.Row]:
        if mode and mode != "all":
            return self.conn.execute(
                "SELECT * FROM scan_history WHERE mode=? ORDER BY id DESC LIMIT ?",
                (mode, limit)
            ).fetchall()
        return self.conn.execute(
            "SELECT * FROM scan_history ORDER BY id DESC LIMIT ?",
            (limit,)
        ).fetchall()

    def get_scan(self, scan_id: int) -> sqlite3.Row | None:
        return self.conn.execute("SELECT * FROM scan_history WHERE id=?", (scan_id,)).fetchone()

    def get_scan_jobs(self, scan_id: int) -> list[sqlite3.Row]:
        # First query by scan_id
        jobs = self.conn.execute(
            "SELECT * FROM jobs WHERE scan_id=? ORDER BY score DESC, fetched_at DESC",
            (scan_id,)
        ).fetchall()
        if jobs:
            return jobs
        # Fallback: check new_job_ids JSON array from scan_history row
        scan = self.get_scan(scan_id)
        if not scan or not scan["new_job_ids"]:
            return []
        try:
            ids = json.loads(scan["new_job_ids"])
            if not ids:
                return []
            placeholders = ",".join("?" for _ in ids)
            return self.conn.execute(
                f"SELECT * FROM jobs WHERE id IN ({placeholders}) ORDER BY score DESC, fetched_at DESC",
                ids
            ).fetchall()
        except Exception:
            return []

    def delete_scan(self, scan_id: int) -> bool:
        cur = self.conn.execute("DELETE FROM scan_history WHERE id=?", (scan_id,))
        self.conn.commit()
        return cur.rowcount > 0

    def clear_scan_history(self) -> int:
        cur = self.conn.execute("DELETE FROM scan_history")
        self.conn.commit()
        return cur.rowcount

    def set(self, job_id: str, **fields):
        cols = ", ".join(f"{k}=?" for k in fields)
        self.conn.execute(f"UPDATE jobs SET {cols} WHERE id=?", (*fields.values(), job_id))
        self.conn.commit()

    def by_status(self, status: str, min_score: int = 0, experience_level: str = "") -> list[sqlite3.Row]:
        if experience_level:
            return self.conn.execute(
                "SELECT * FROM jobs WHERE status=? AND score>=? AND experience_level=? ORDER BY score DESC",
                (status, min_score, experience_level)).fetchall()
        return self.conn.execute(
            "SELECT * FROM jobs WHERE status=? AND score>=? ORDER BY score DESC",
            (status, min_score)).fetchall()

    def get(self, job_id: str) -> sqlite3.Row | None:
        return self.conn.execute("SELECT * FROM jobs WHERE id LIKE ?", (job_id + "%",)).fetchone()

    def all(self) -> list[sqlite3.Row]:
        return self.conn.execute("SELECT * FROM jobs ORDER BY score DESC").fetchall()

    def export_xlsx(self, path: Path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Jobs"
        cols = ["id", "status", "score", "title", "company", "location", "salary",
                "source", "posted_at", "url", "applied_at", "notes",
                "experience_level", "company_careers_url"]
        ws.append([c.upper() for c in cols])
        for row in self.all():
            ws.append([row[c] for c in cols])
        for i, w in enumerate([18, 12, 7, 40, 25, 22, 15, 12, 20, 60, 20, 30], 1):
            ws.column_dimensions[chr(64 + i)].width = w
        wb.save(path)

